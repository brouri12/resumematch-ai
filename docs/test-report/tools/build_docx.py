"""Convert a Markdown report (default TECHNICAL_REPORT.md) into .docx (headings, tables, images, code, lists).

Usage: python build_docx.py [source.md]  -> writes source.docx next to it.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "TECHNICAL_REPORT.md"
ROOT = SRC.parent
DEST = SRC.with_suffix(".docx")

IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
FOREST = RGBColor(0x1F, 0x4D, 0x3E)


def shade(cell, hex_color):
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), hex_color)
    props.append(fill)


def add_inline(paragraph, text, size=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if bold:
            run.bold = True
        if size:
            run.font.size = Pt(size)


def add_image(container, path, width_cm):
    image = ROOT / path
    if not image.exists():
        return False
    paragraph = container.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image), width=Cm(width_cm))
    return True


def parse_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    header, body = rows[0], [r for r in rows[1:] if not set("".join(r)) <= set("-: ")]
    has_images = any(IMAGE.search(c) for r in [header, *body] for c in r)
    if has_images:
        body = [r for r in [header, *body] if any(c for c in r)]
        header = None
    cols = max(len(r) for r in ([header] if header else []) + body)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    font_size = 8 if cols >= 7 else 9
    for index, row in enumerate(([header] if header else []) + body):
        cells = table.add_row().cells
        for col in range(cols):
            text = row[col] if col < len(row) else ""
            cell = cells[col]
            cell.paragraphs[0].text = ""
            images = IMAGE.findall(text)
            if images:
                for _alt, path in images:
                    pic = cell.paragraphs[0]
                    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pic.add_run().add_picture(str(ROOT / path), width=Cm(7.6))
                caption = IMAGE.sub("", text).strip()
                if caption:
                    add_inline(cell.add_paragraph(), caption, size=8)
            else:
                add_inline(cell.paragraphs[0], text, size=font_size, bold=header is not None and index == 0)
            if header is not None and index == 0:
                shade(cell, "DCEFE6")
    doc.add_paragraph()


def main():
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(1.6)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for level in (1, 2, 3):
        styles[f"Heading {level}"].font.color.rgb = FOREST

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            paragraph = doc.add_paragraph()
            run = paragraph.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            paragraph.paragraph_format.left_indent = Cm(0.4)
            shade_par = OxmlElement("w:shd")
            shade_par.set(qn("w:val"), "clear")
            shade_par.set(qn("w:fill"), "F2F4F3")
            paragraph._p.get_or_add_pPr().append(shade_par)
        elif stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        elif stripped.startswith("# "):
            title = doc.add_heading(stripped[2:], level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].replace("`", ""), level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].replace("`", ""), level=1)
        elif stripped == "---":
            pass
        elif IMAGE.fullmatch(stripped):
            alt, path = IMAGE.fullmatch(stripped).groups()
            if add_image(doc, path, 16.5):
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(f"Figure — {alt}")
                run.italic = True
                run.font.size = Pt(9)
        elif stripped.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            add_inline(paragraph, stripped[2:])
            for run in paragraph.runs:
                run.italic = True
        elif stripped.startswith("- "):
            add_inline(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif stripped:
            add_inline(doc.add_paragraph(), stripped)
        i += 1

    doc.save(DEST)
    print("wrote", DEST)


if __name__ == "__main__":
    main()
