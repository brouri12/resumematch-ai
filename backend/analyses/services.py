"""Analysis orchestration: scoring, recommendations, cover letters."""

from __future__ import annotations

import logging
import threading
from datetime import timedelta
from io import BytesIO
from typing import Callable

from django.conf import settings
from django.db import close_old_connections
from django.utils import timezone

from cvs.models import CV
from services.ai_service import get_ai_service
from services.ats_service import collect_pdf_layout, evaluate_ats
from services.scoring import canonicalize_skills, skill_overlap_score

from .models import Analysis, AnalysisJob, CoverLetter, JobOffer

logger = logging.getLogger(__name__)

ProgressFn = Callable[[int, str], None]


def _noop_progress(_pct: int, _step: str) -> None:
    return None


def run_analysis(
    *,
    user,
    cv: CV,
    job_offer_text: str,
    title: str = "",
    company: str = "",
    level: str = Analysis.Level.JUNIOR,
    language: str = "fr",
    progress: ProgressFn = _noop_progress,
) -> Analysis:
    ai = get_ai_service()

    progress(10, "cv")
    if not cv.parsed_data and cv.extracted_text:
        cv.parsed_data = ai.extract_cv_skills(cv.extracted_text).model_dump()
        cv.save(update_fields=["parsed_data"])

    progress(25, "job")
    job_parsed = ai.analyze_job_offer(job_offer_text)
    resolved_title = title.strip() or job_parsed.title or "Offre sans titre"
    resolved_company = company.strip() or job_parsed.company or ""

    job_offer = JobOffer.objects.create(
        user=user,
        title=resolved_title,
        company=resolved_company,
        raw_text=job_offer_text,
        parsed_data=job_parsed.model_dump(),
    )

    cv_skills = list(cv.parsed_data.get("technical_skills") or []) + list(
        cv.parsed_data.get("soft_skills") or []
    )
    required = list(job_parsed.required_skills or [])
    nice = list(job_parsed.nice_to_have_skills or [])

    skills_score, present, missing = skill_overlap_score(cv_skills, required)

    # Nice-to-have bonus (small)
    nice_present = 0
    if nice:
        _, nice_hit, _ = skill_overlap_score(cv_skills, nice)
        nice_present = len(nice_hit)

    progress(45, "semantic")
    semantic = ai.evaluate_semantic_fit(
        cv.extracted_text or "",
        job_offer_text,
        level=level,  # type: ignore[arg-type]
        present_skills=present,
        missing_skills=missing,
    )

    w_skills = settings.SCORE_WEIGHT_SKILLS
    w_semantic = settings.SCORE_WEIGHT_SEMANTIC
    level_mult = settings.SCORE_LEVEL_MULTIPLIERS.get(level, 1.0)

    # Renormalize weights after level multiplier on semantic part
    adj_semantic = w_semantic * level_mult
    total_w = w_skills + adj_semantic
    skills_contrib = (w_skills / total_w) * skills_score
    semantic_contrib = (adj_semantic / total_w) * semantic.semantic_score
    nice_bonus = min(5.0, nice_present * 1.5)
    final_score = int(round(min(100.0, skills_contrib + semantic_contrib + nice_bonus)))

    progress(70, "recommendations")
    recs = ai.generate_recommendations(
        cv.extracted_text or "",
        job_offer_text,
        level=level,  # type: ignore[arg-type]
        missing_skills=missing,
        detailed=False,
        language=language,  # type: ignore[arg-type]
    )
    progress(90, "saving")

    breakdown = {
        "skills_score": round(skills_score, 1),
        "semantic_score": semantic.semantic_score,
        "skills_weight": round(w_skills / total_w, 3),
        "semantic_weight": round(adj_semantic / total_w, 3),
        "skills_contribution": round(skills_contrib, 1),
        "semantic_contribution": round(semantic_contrib, 1),
        "nice_to_have_bonus": round(nice_bonus, 1),
        "level": level,
        "level_semantic_multiplier": level_mult,
        "rationale": semantic.rationale,
        "experience_fit": semantic.experience_fit,
        "present_count": len(present),
        "missing_count": len(missing),
        "required_count": len(canonicalize_skills(required)),
    }

    analysis = Analysis.objects.create(
        user=user,
        cv=cv,
        job_offer=job_offer,
        level=level,
        score=final_score,
        present_skills=present,
        missing_skills=missing,
        recommendations=recs.recommandations,
        score_breakdown=breakdown,
        status=Analysis.Status.TO_APPLY,
        language=language,
        ai_mode=ai.mode,
    )
    progress(100, "done")
    return analysis


def analyses_today(user) -> int:
    since = timezone.now() - timedelta(days=1)
    return Analysis.objects.filter(user=user, created_at__gte=since).count()


def effective_daily_quota() -> int:
    from adminpanel.models import SiteSettings

    return SiteSettings.load().effective_daily_quota


def quota_remaining(user, extra_pending: int = 0) -> int:
    limit = effective_daily_quota()
    if limit <= 0:
        return 10**6
    pending = AnalysisJob.objects.filter(
        user=user, state__in=[AnalysisJob.State.PENDING, AnalysisJob.State.RUNNING]
    ).count()
    return max(0, limit - analyses_today(user) - pending - extra_pending)


def start_analysis_job(*, user, cv: CV, **kwargs) -> AnalysisJob:
    job = AnalysisJob.objects.create(user=user)

    def update(pct: int, step: str) -> None:
        AnalysisJob.objects.filter(pk=job.pk).update(
            state=AnalysisJob.State.RUNNING, progress=pct, step=step, updated_at=timezone.now()
        )

    in_thread = not settings.ANALYSIS_JOBS_SYNC

    def worker() -> None:
        if in_thread:
            close_old_connections()
        try:
            analysis = run_analysis(user=user, cv=cv, progress=update, **kwargs)
            AnalysisJob.objects.filter(pk=job.pk).update(
                state=AnalysisJob.State.DONE,
                progress=100,
                step="done",
                analysis=analysis,
                updated_at=timezone.now(),
            )
        except Exception as exc:  # noqa: BLE001
            logger.exception("Analysis job %s failed", job.pk)
            AnalysisJob.objects.filter(pk=job.pk).update(
                state=AnalysisJob.State.FAILED,
                error=str(exc)[:500] or "Erreur inconnue",
                updated_at=timezone.now(),
            )
        finally:
            if in_thread:
                close_old_connections()

    if in_thread:
        threading.Thread(target=worker, daemon=True, name=f"analysis-job-{job.pk}").start()
    else:
        worker()
        job.refresh_from_db()
    return job


def _extras_key(kind: str, language: str) -> str:
    return f"{kind}:{language}"


def get_or_generate_extra(analysis: Analysis, kind: str, language: str, refresh: bool = False) -> dict:
    """Generate (or return cached) interview_prep / cv_rewrite / learning_plan / ats content."""
    key = _extras_key(kind, language)
    extras = dict(analysis.extras or {})
    if not refresh and key in extras:
        return extras[key]

    ai = get_ai_service()
    cv_text = analysis.cv.extracted_text or ""
    job_text = analysis.job_offer.raw_text
    missing = analysis.missing_skills or []

    if kind == "interview_prep":
        data = ai.generate_interview_prep(
            cv_text, job_text, analysis.level, missing, language  # type: ignore[arg-type]
        ).model_dump()
    elif kind == "cv_rewrite":
        data = ai.rewrite_cv(
            cv_text, job_text, analysis.level, missing, language  # type: ignore[arg-type]
        ).model_dump()
    elif kind == "learning_plan":
        data = ai.generate_learning_plan(
            missing, analysis.level, analysis.job_offer.title, language  # type: ignore[arg-type]
        ).model_dump()
    elif kind == "ats":
        layout = {}
        try:
            layout = collect_pdf_layout(analysis.cv.file.path)
        except Exception:  # noqa: BLE001
            logger.warning("CV file unavailable for ATS layout", exc_info=True)
        required = list((analysis.job_offer.parsed_data or {}).get("required_skills") or [])
        data = evaluate_ats(cv_text, layout, required, language)
    else:
        raise ValueError(f"Unknown extra kind: {kind}")

    if kind != "ats":
        data["ai_mode"] = ai.mode
    extras[key] = data
    analysis.extras = extras
    analysis.save(update_fields=["extras"])
    return data


def build_report_docx(analysis: Analysis, language: str = "fr") -> bytes:
    from docx import Document

    en = language == "en"
    t = (lambda fr, eng: eng if en else fr)
    breakdown = analysis.score_breakdown or {}
    doc = Document()
    doc.add_heading(t("Rapport d'analyse – ResumeMatch AI", "Analysis report – ResumeMatch AI"), 0)
    doc.add_paragraph(
        f"{analysis.job_offer.title} · {analysis.job_offer.company or '—'} · "
        f"{analysis.created_at:%d/%m/%Y}"
    )
    doc.add_heading(t(f"Score : {analysis.score}/100", f"Score: {analysis.score}/100"), 1)
    doc.add_paragraph(
        t("Compétences", "Skills") + f" : {breakdown.get('skills_score', '—')} · "
        + t("Sémantique", "Semantic") + f" : {breakdown.get('semantic_score', '—')}"
    )
    if breakdown.get("rationale"):
        doc.add_paragraph(str(breakdown["rationale"]))

    def bullets(title: str, items: list) -> None:
        doc.add_heading(title, 2)
        if not items:
            doc.add_paragraph("—")
        for item in items:
            doc.add_paragraph(str(item), style="List Bullet")

    bullets(t("Compétences présentes", "Matched skills"), analysis.present_skills or [])
    bullets(t("Compétences manquantes", "Missing skills"), analysis.missing_skills or [])
    bullets(t("Recommandations", "Recommendations"), analysis.recommendations or [])

    extras = analysis.extras or {}
    ats = extras.get(_extras_key("ats", language))
    if ats:
        doc.add_heading(t(f"Compatibilité ATS : {ats['score']}/100", f"ATS compatibility: {ats['score']}/100"), 2)
        for check in ats.get("checks", []):
            doc.add_paragraph(f"[{check['status'].upper()}] {check['label']} – {check['detail']}", style="List Bullet")

    rewrite = extras.get(_extras_key("cv_rewrite", language))
    if rewrite:
        doc.add_heading(t("Réécriture du CV", "CV rewrite"), 2)
        for item in rewrite.get("items", []):
            if item.get("before"):
                doc.add_paragraph(t("Avant : ", "Before: ") + item["before"])
            doc.add_paragraph(t("Après : ", "After: ") + item["after"])

    interview = extras.get(_extras_key("interview_prep", language))
    if interview:
        doc.add_heading(t("Préparation d'entretien", "Interview preparation"), 2)
        for q in interview.get("questions", []):
            doc.add_paragraph(q["question"], style="List Number")
            if q.get("answer_tips"):
                doc.add_paragraph(q["answer_tips"])

    plan = extras.get(_extras_key("learning_plan", language))
    if plan:
        doc.add_heading(t("Plan d'apprentissage", "Learning plan"), 2)
        for step in plan.get("steps", []):
            doc.add_paragraph(
                f"{step['skill']} – {step['priority']} – ~{step['estimated_weeks']} "
                + t("semaine(s)", "week(s)"),
                style="List Bullet",
            )
            if step.get("project_idea"):
                doc.add_paragraph(step["project_idea"])

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def improve_cv_suggestions(analysis: Analysis, language: str = "fr") -> dict:
    ai = get_ai_service()
    result = ai.generate_recommendations(
        analysis.cv.extracted_text or "",
        analysis.job_offer.raw_text,
        level=analysis.level,  # type: ignore[arg-type]
        missing_skills=analysis.missing_skills or [],
        detailed=True,
        language=language,  # type: ignore[arg-type]
    )
    return {
        "recommandations": result.recommandations,
        "detailed_improvements": result.detailed_improvements,
        "niveau": analysis.level,
        "analysis_id": analysis.id,
        "ai_mode": ai.mode,
    }


def create_cover_letter(analysis: Analysis, tone: str, language: str = "fr") -> CoverLetter:
    ai = get_ai_service()
    user = analysis.user
    name = (user.get_full_name() or user.username).strip()
    result = ai.generate_cover_letter(
        analysis.cv.extracted_text or "",
        analysis.job_offer.raw_text,
        level=analysis.level,  # type: ignore[arg-type]
        tone=tone,  # type: ignore[arg-type]
        candidate_name=name,
        job_title=analysis.job_offer.title,
        company=analysis.job_offer.company,
        language=language,  # type: ignore[arg-type]
    )
    return CoverLetter.objects.create(
        analysis=analysis,
        tone=tone,
        content=result.content,
    )
